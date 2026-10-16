"""
Reporte automatico de ventas.

Une todos los archivos Excel (.xlsx) y CSV de una carpeta, limpia los datos
y genera un unico Excel con resumen, totales por producto, vendedor y mes.

Uso:
    python reporte_ventas.py                      # usa la carpeta "datos"
    python reporte_ventas.py mi_carpeta -o reporte.xlsx
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS = ["Fecha", "Producto", "Cantidad", "Precio_Unitario", "Vendedor"]


def leer_archivos(carpeta: Path) -> pd.DataFrame:
    """Lee todos los .xlsx y .csv de la carpeta y los junta en una tabla."""
    archivos = sorted(
        p for p in carpeta.iterdir()
        if p.suffix.lower() in {".xlsx", ".csv"} and not p.name.startswith("~$")
    )
    if not archivos:
        sys.exit(f"No encontre archivos .xlsx ni .csv en '{carpeta}'.")

    tablas = []
    for archivo in archivos:
        if archivo.suffix.lower() == ".csv":
            tabla = pd.read_csv(archivo)
        else:
            tabla = pd.read_excel(archivo)
        faltan = [c for c in COLUMNAS if c not in tabla.columns]
        if faltan:
            print(f"Aviso: '{archivo.name}' se omitio, le faltan columnas: {faltan}")
            continue
        tabla["Archivo_Origen"] = archivo.name
        tablas.append(tabla)
        print(f"Leido: {archivo.name} ({len(tabla)} filas)")

    if not tablas:
        sys.exit("Ningun archivo tenia las columnas necesarias.")
    return pd.concat(tablas, ignore_index=True)


def limpiar(df: pd.DataFrame, quitar_duplicados: bool = False) -> pd.DataFrame:
    """Quita filas vacias, espacios sobrantes y datos invalidos."""
    antes = len(df)
    df = df.dropna(how="all").copy()

    df["Producto"] = df["Producto"].astype(str).str.strip().str.title()
    df["Vendedor"] = df["Vendedor"].astype(str).str.strip().str.title()
    df["Fecha"] = pd.to_datetime(df["Fecha"], errors="coerce")
    df["Cantidad"] = pd.to_numeric(df["Cantidad"], errors="coerce")
    df["Precio_Unitario"] = pd.to_numeric(df["Precio_Unitario"], errors="coerce")

    df = df.dropna(subset=["Fecha", "Producto", "Cantidad", "Precio_Unitario"])
    df = df[(df["Cantidad"] > 0) & (df["Precio_Unitario"] >= 0)]
    if quitar_duplicados:
        df = df.drop_duplicates(subset=COLUMNAS)

    df["Total"] = (df["Cantidad"] * df["Precio_Unitario"]).round(2)
    df["Mes"] = df["Fecha"].dt.to_period("M").astype(str)
    print(f"Limpieza: {antes} filas leidas, {len(df)} validas.")
    return df.sort_values("Fecha").reset_index(drop=True)


def crear_resumen(df: pd.DataFrame) -> pd.DataFrame:
    mejor_producto = df.groupby("Producto")["Total"].sum().idxmax()
    mejor_vendedor = df.groupby("Vendedor")["Total"].sum().idxmax()
    filas = [
        ("Total vendido", round(df["Total"].sum(), 2)),
        ("Numero de ventas", len(df)),
        ("Unidades vendidas", int(df["Cantidad"].sum())),
        ("Venta promedio", round(df["Total"].mean(), 2)),
        ("Producto mas vendido (en dinero)", mejor_producto),
        ("Mejor vendedor (en dinero)", mejor_vendedor),
        ("Desde", df["Fecha"].min().strftime("%Y-%m-%d")),
        ("Hasta", df["Fecha"].max().strftime("%Y-%m-%d")),
    ]
    return pd.DataFrame(filas, columns=["Indicador", "Valor"])


def agrupar(df: pd.DataFrame, columna: str) -> pd.DataFrame:
    tabla = (
        df.groupby(columna)
        .agg(Ventas=("Total", "count"), Unidades=("Cantidad", "sum"), Total=("Total", "sum"))
        .sort_values("Total", ascending=False)
        .reset_index()
    )
    tabla["Total"] = tabla["Total"].round(2)
    return tabla


def dar_formato(escritor, hojas_con_grafico: dict) -> None:
    """Encabezados en negrita, columnas ajustadas y graficos de barras."""
    relleno = PatternFill("solid", fgColor="1F4E78")
    for nombre, hoja in escritor.sheets.items():
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center")
        for i, columna in enumerate(hoja.columns, start=1):
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
            hoja.column_dimensions[get_column_letter(i)].width = min(ancho + 4, 45)
        hoja.freeze_panes = "A2"

    for nombre, (col_total, titulo) in hojas_con_grafico.items():
        hoja = escritor.sheets[nombre]
        grafico = BarChart()
        grafico.title = titulo
        grafico.legend = None
        datos = Reference(hoja, min_col=col_total, min_row=1, max_row=hoja.max_row)
        etiquetas = Reference(hoja, min_col=1, min_row=2, max_row=hoja.max_row)
        grafico.add_data(datos, titles_from_data=True)
        grafico.set_categories(etiquetas)
        grafico.height, grafico.width = 8, 16
        hoja.add_chart(grafico, "G2")


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera un reporte de ventas en Excel.")
    parser.add_argument("carpeta", nargs="?", default="datos", help="carpeta con los archivos")
    parser.add_argument("-o", "--salida", default="reporte_ventas.xlsx", help="archivo de salida")
    parser.add_argument("--quitar-duplicados", action="store_true",
                        help="elimina filas totalmente repetidas (util si hay archivos duplicados)")
    args = parser.parse_args()

    carpeta = Path(args.carpeta)
    if not carpeta.is_dir():
        sys.exit(f"La carpeta '{carpeta}' no existe.")

    df = limpiar(leer_archivos(carpeta), args.quitar_duplicados)
    if df.empty:
        sys.exit("No quedaron datos validos despues de la limpieza.")

    por_producto = agrupar(df, "Producto")
    por_vendedor = agrupar(df, "Vendedor")
    por_mes = agrupar(df, "Mes").sort_values("Mes")

    with pd.ExcelWriter(args.salida, engine="openpyxl") as escritor:
        crear_resumen(df).to_excel(escritor, sheet_name="Resumen", index=False)
        por_producto.to_excel(escritor, sheet_name="Por producto", index=False)
        por_vendedor.to_excel(escritor, sheet_name="Por vendedor", index=False)
        por_mes.to_excel(escritor, sheet_name="Por mes", index=False)
        df.assign(Fecha=df["Fecha"].dt.date).to_excel(escritor, sheet_name="Datos", index=False)
        dar_formato(
            escritor,
            {
                "Por producto": (4, "Total vendido por producto"),
                "Por vendedor": (4, "Total vendido por vendedor"),
                "Por mes": (4, "Total vendido por mes"),
            },
        )

    print(f"\nReporte creado: {args.salida}")


if __name__ == "__main__":
    main()
